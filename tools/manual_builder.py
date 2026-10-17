# -*- coding: utf-8 -*-
"""Build the organizer operation manual as an editable .docx file."""

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


BODY_FONT = "宋体"
HEAD_FONT = "微软雅黑"
BODY_SIZE = 10.5
TABLE_SIZE = 9
MAIN_COLOR = RGBColor(0x1F, 0x3A, 0x5F)
ACCENT_COLOR = RGBColor(0xB0, 0x6A, 0x1E)
GRAY = RGBColor(0x66, 0x6F, 0x76)


def _set_run_font(run, name=BODY_FONT, size=BODY_SIZE, bold=None, color=None, italic=None):
    run.font.name = name
    run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), name)


def _shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _set_table_borders(table, color="B7C4D1", size="4"):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tbl_pr.append(borders)


def _repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def _keep_row_together(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def _add_field(paragraph, instruction, placeholder=""):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    run._r.append(text)
    run._r.append(fld_end)
    return run


class Manual:
    def __init__(self, doc):
        self.doc = doc
        self._step_no = 0

    # ---------- text ----------
    def _rich(self, paragraph, text, size=BODY_SIZE, color=None, bold=False, italic=False, font=BODY_FONT):
        parts = text.split("**")
        for index, part in enumerate(parts):
            if not part:
                continue
            run = paragraph.add_run(part)
            _set_run_font(
                run,
                name=font,
                size=size,
                bold=(bold or index % 2 == 1),
                color=color,
                italic=italic,
            )

    def h1(self, text):
        p = self.doc.add_paragraph(style="Heading 1")
        run = p.add_run(text)
        _set_run_font(run, HEAD_FONT, 18, bold=True, color=MAIN_COLOR)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(14)
        self._step_no = 0

    def h2(self, text):
        p = self.doc.add_paragraph(style="Heading 2")
        run = p.add_run(text)
        _set_run_font(run, HEAD_FONT, 14, bold=True, color=MAIN_COLOR)
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(7)
        self._step_no = 0

    def h3(self, text):
        p = self.doc.add_paragraph(style="Heading 3")
        run = p.add_run(text)
        _set_run_font(run, HEAD_FONT, 11.5, bold=True, color=RGBColor(0x2B, 0x2B, 0x2B))
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        self._step_no = 0

    def p(self, text, gray=False, size=BODY_SIZE):
        p = self.doc.add_paragraph()
        self._rich(p, text, size=size, color=GRAY if gray else None)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.35
        return p

    def meta(self, text):
        p = self.doc.add_paragraph()
        self._rich(p, text, size=9, color=GRAY, italic=True)
        p.paragraph_format.space_after = Pt(5)
        return p

    def bullet(self, text, level=0):
        p = self.doc.add_paragraph()
        indent = 0.65 + level * 0.55
        p.paragraph_format.left_indent = Cm(indent)
        p.paragraph_format.first_line_indent = Cm(-0.45)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.3
        run = p.add_run("•  ")
        _set_run_font(run, BODY_FONT, BODY_SIZE, color=ACCENT_COLOR, bold=True)
        self._rich(p, text)
        return p

    def steps(self, items):
        for item in items:
            self._step_no += 1
            p = self.doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.9)
            p.paragraph_format.first_line_indent = Cm(-0.9)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.3
            run = p.add_run(f"{self._step_no}. ")
            _set_run_font(run, BODY_FONT, BODY_SIZE, bold=True, color=MAIN_COLOR)
            self._rich(p, item)

    def reset_steps(self):
        self._step_no = 0

    # ---------- blocks ----------
    def table(self, headers, rows, widths=None, font_size=TABLE_SIZE, header_fill="2F4B7C"):
        table = self.doc.add_table(rows=1, cols=len(headers))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        _set_table_borders(table)
        header_row = table.rows[0]
        _repeat_table_header(header_row)
        _keep_row_together(header_row)
        for i, text in enumerate(headers):
            cell = header_row.cells[i]
            cell.text = ""
            _shade_cell(cell, header_fill)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(text)
            _set_run_font(run, HEAD_FONT, font_size, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        for row_data in rows:
            new_row = table.add_row()
            _keep_row_together(new_row)
            cells = new_row.cells
            for i, text in enumerate(row_data):
                cell = cells[i]
                cell.text = ""
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                lines = str(text).split("\n")
                for index, line in enumerate(lines):
                    p = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.line_spacing = 1.15
                    self._rich(p, line, size=font_size)
        if widths:
            for row in table.rows:
                for i, width in enumerate(widths):
                    row.cells[i].width = Cm(width)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)
        return table

    def callout(self, label, text, kind="tip"):
        palettes = {
            "tip": ("E8F1FB", "2F6FA8", "提示"),
            "note": ("FFF4E2", "B06A1E", "注意"),
            "platform": ("EFEFEF", "666666", "平台自办差异"),
            "shot": ("EAF4EA", "3C7A45", "截图位"),
        }
        fill, border, default_label = palettes.get(kind, palettes["tip"])
        table = self.doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        _set_table_borders(table, color=border, size="6")
        cell = table.cell(0, 0)
        _shade_cell(cell, fill)
        cell.text = ""
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(f"【{label or default_label}】")
        _set_run_font(run, HEAD_FONT, 9.5, bold=True, color=RGBColor.from_string(border))
        self._rich(p, text, size=9.5)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)
        return table

    def screenshot(self, code, title, requirement="", annotation=""):
        table = self.doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        _set_table_borders(table, color="7FA7D9", size="6")
        cell = table.cell(0, 0)
        _shade_cell(cell, "F2F7FD")
        cell.text = ""
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(f"[截图位 {code}] ")
        _set_run_font(run, HEAD_FONT, 9.5, bold=True, color=RGBColor(0x2F, 0x6F, 0xA8))
        self._rich(p, title, size=9.5, bold=True, color=RGBColor(0x2F, 0x6F, 0xA8))
        if requirement:
            p2 = cell.add_paragraph()
            p2.paragraph_format.space_after = Pt(1)
            self._rich(p2, f"画面要求：{requirement}", size=9, color=GRAY)
        if annotation:
            p3 = cell.add_paragraph()
            p3.paragraph_format.space_after = Pt(1)
            self._rich(p3, f"标注建议：{annotation}", size=9, color=GRAY)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)
        return table

    def page_break(self):
        self.doc.add_page_break()

    # ---------- cover & front matter ----------
    def cover(self):
        for _ in range(3):
            self.doc.add_paragraph()
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run("啤酒赛事平台")
        _set_run_font(run, HEAD_FONT, 22, bold=True, color=ACCENT_COLOR)
        p.paragraph_format.space_after = Pt(12)

        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run("主办方全流程操作手册")
        _set_run_font(run, HEAD_FONT, 32, bold=True, color=MAIN_COLOR)
        p.paragraph_format.space_after = Pt(10)

        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run("SaaS 版 · V1.0")
        _set_run_font(run, HEAD_FONT, 15, bold=True, color=GRAY)
        p.paragraph_format.space_after = Pt(30)

        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run("从入驻申请、赛前配置、报名收样、现场评审到结果发布与归档")
        _set_run_font(run, BODY_FONT, 12, color=GRAY)
        p.paragraph_format.space_after = Pt(6)

        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run("适用对象：第三方赛事主办方主管理员 / 子管理员 / 赛事执行人员")
        _set_run_font(run, BODY_FONT, 11, color=GRAY)
        p.paragraph_format.space_after = Pt(60)

        for _ in range(4):
            self.doc.add_paragraph()

        info = self.doc.add_table(rows=4, cols=2)
        info.alignment = WD_TABLE_ALIGNMENT.CENTER
        _set_table_borders(info, color="D6DEE7", size="4")
        info_rows = [
            ("文档版本", "V1.0（初稿）"),
            ("编制日期", "2026 年 10 月"),
            ("适用系统", "啤酒赛事平台 · 主办方后台 / 厂商端 / 裁判 H5"),
            ("文档状态", "正文完成，截图位待替换"),
        ]
        for i, (key, value) in enumerate(info_rows):
            cells = info.rows[i].cells
            cells[0].text = ""
            cells[1].text = ""
            cells[0].width = Cm(4.2)
            cells[1].width = Cm(9.0)
            _shade_cell(cells[0], "F2F5F8")
            r1 = cells[0].paragraphs[0].add_run(key)
            _set_run_font(r1, HEAD_FONT, 10, bold=True, color=MAIN_COLOR)
            r2 = cells[1].paragraphs[0].add_run(value)
            _set_run_font(r2, BODY_FONT, 10)

    def toc(self):
        self.doc.add_page_break()
        p = self.doc.add_paragraph()
        run = p.add_run("目录")
        _set_run_font(run, HEAD_FONT, 18, bold=True, color=MAIN_COLOR)
        p.paragraph_format.space_after = Pt(12)

        p = self.doc.add_paragraph()
        run = p.add_run()
        fld_begin = OxmlElement("w:fldChar")
        fld_begin.set(qn("w:fldCharType"), "begin")
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = 'TOC \\o "1-2" \\h \\z \\u'
        fld_sep = OxmlElement("w:fldChar")
        fld_sep.set(qn("w:fldCharType"), "separate")
        text = OxmlElement("w:t")
        text.text = "打开文档后，右键点击此处并选择“更新域”，即可生成目录。"
        fld_end = OxmlElement("w:fldChar")
        fld_end.set(qn("w:fldCharType"), "end")
        run._r.append(fld_begin)
        run._r.append(instr)
        run._r.append(fld_sep)
        run._r.append(text)
        run._r.append(fld_end)
        _set_run_font(run, BODY_FONT, 10, color=GRAY)
        self.doc.add_page_break()


def create_document():
    doc = Document()

    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.0)
    section.header_distance = Cm(1.1)
    section.footer_distance = Cm(1.0)
    section.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(BODY_SIZE)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.space_after = Pt(5)

    for style_name, font, size, color in (
        ("Heading 1", HEAD_FONT, 18, MAIN_COLOR),
        ("Heading 2", HEAD_FONT, 14, MAIN_COLOR),
        ("Heading 3", HEAD_FONT, 11.5, RGBColor(0x2B, 0x2B, 0x2B)),
    ):
        style = doc.styles[style_name]
        style.font.name = font
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.element.rPr.rFonts.set(qn("w:eastAsia"), font)
        style.paragraph_format.keep_with_next = True

    header_p = section.header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header_p.add_run("啤酒赛事平台 · 主办方全流程操作手册（SaaS 版）")
    _set_run_font(run, HEAD_FONT, 8.5, color=GRAY)
    p_pr = header_p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), "C9D2DB")
    borders.append(bottom)
    p_pr.append(borders)

    footer_p = section.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer_p.add_run("第 ")
    _set_run_font(run, BODY_FONT, 9, color=GRAY)
    _add_field(footer_p, "PAGE", "1")
    run = footer_p.add_run(" 页 / 共 ")
    _set_run_font(run, BODY_FONT, 9, color=GRAY)
    _add_field(footer_p, "NUMPAGES", "1")
    run = footer_p.add_run(" 页")
    _set_run_font(run, BODY_FONT, 9, color=GRAY)

    settings = doc.settings.element
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    settings.append(update_fields)

    return doc


def save(doc, path):
    doc.save(path)
