from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import uuid
from collections import OrderedDict, defaultdict
from decimal import Decimal, InvalidOperation
from pathlib import Path

import openpyxl


COMPETITION_ID = 1070
SOURCE_VERSION = "REGISTRATION_FORM_20260801"
INPUT_PATH = Path(r"C:\Users\a1508\Desktop\首届中国拉格大赛_报名表.xlsx")

CATEGORY_BY_TEXT = OrderedDict(
    [
        ("浅色拉格组", 1593),
        ("深色拉格组", 1594),
        ("创意拉格组", 1595),
    ]
)

# The user explicitly requested these merges.
BRAND_RENAMES = {
    "狐朋": "嘉樂酿造",
    "桂南嘉樂": "嘉樂酿造",
    "水门头": "狂歌",
}

# 后续补充表中的手机号。按作品编号绑定，避免 Excel 删除未付款行后行号变化。
PHONE_OVERRIDES_BY_WORK_CODE = {
    "503137760384": "18622027553",  # 苍神酿造／壹苇渡江
    "685277144062": "13963925662",  # 脸谱酿造／夏至
}

STYLE_ABV_SWAPPED_WORK_CODES = {"1545298644322"}
EMPTY_SPECIAL_WORK_CODES = {"2325330145824", "9796559967045"}


def clean(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def excel_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def sql_text(value: str | None) -> str:
    if value is None:
        return "NULL"
    if value == "":
        return "''"
    raw = str(value).encode("utf-8").hex()
    return f"CONVERT(0x{raw} USING utf8mb4)"


def sql_number(value: Decimal | int | float | str | None) -> str:
    if value is None:
        return "NULL"
    return str(value)


def parse_abv(value: str) -> Decimal:
    match = re.search(r"\d+(?:\.\d+)?", value)
    if not match:
        raise ValueError(f"无法解析 ABV: {value!r}")
    try:
        abv = Decimal(match.group(0)).quantize(Decimal("0.01"))
    except InvalidOperation as exc:
        raise ValueError(f"无法解析 ABV: {value!r}") from exc
    if abv < 0 or abv > 99.99:
        raise ValueError(f"ABV 超出数据库范围: {value!r}")
    return abv


def source_work_code(value) -> str:
    code = excel_text(value)
    if not code.isdigit():
        raise ValueError(f"作品编号不是纯数字: {code!r}")
    if len(code) < 5:
        raise ValueError(f"作品编号不足五位: {code!r}")
    return code


def final_brand(raw_brand: str) -> str:
    return BRAND_RENAMES.get(raw_brand, raw_brand)


def category_id(group: str) -> int:
    for marker, category in CATEGORY_BY_TEXT.items():
        if marker in group:
            return category
    raise ValueError(f"无法匹配比赛组别: {group!r}")


def stable_uuid(source_row: int, work_code: str) -> str:
    return str(uuid.uuid5(uuid.NAMESPACE_URL, f"beer-competition:excel:{COMPETITION_ID}:{source_row}:{work_code}"))


def stable_token(source_row: int, work_code: str) -> str:
    digest = hashlib.sha256(f"{COMPETITION_ID}:{source_row}:{work_code}".encode("ascii")).hexdigest().upper()
    return "Q" + digest[:63]


def load_rows(path: Path) -> list[dict]:
    workbook = openpyxl.load_workbook(path, data_only=False)
    sheet = workbook.active
    rows: list[dict] = []
    for row_no in range(3, sheet.max_row + 1):
        if clean(sheet.cell(row_no, 2).value) != "已付款":
            continue

        work_code = source_work_code(sheet.cell(row_no, 13).value)
        raw_style = clean(sheet.cell(row_no, 9).value)
        raw_abv = clean(sheet.cell(row_no, 10).value)
        if work_code in STYLE_ABV_SWAPPED_WORK_CODES:
            raw_style, raw_abv = raw_abv, raw_style

        raw_special = clean(sheet.cell(row_no, 11).value)
        if work_code in EMPTY_SPECIAL_WORK_CODES or not raw_special:
            raw_special = "无"

        phone = clean(sheet.cell(row_no, 5).value)
        raw_brand = clean(sheet.cell(row_no, 6).value)
        if raw_brand == "悠航鲜啤":
            phone = "13718007058"
        phone = PHONE_OVERRIDES_BY_WORK_CODE.get(work_code, phone)

        brand = final_brand(raw_brand)
        group = clean(sheet.cell(row_no, 7).value)
        short_code = work_code[-5:]

        rows.append(
            {
                "source_row": row_no,
                "contact_name": clean(sheet.cell(row_no, 3).value),
                "wechat": clean(sheet.cell(row_no, 4).value),
                "phone": phone,
                "raw_brand": raw_brand,
                "brand": brand,
                "group": group,
                "category_id": category_id(group),
                "name": clean(sheet.cell(row_no, 8).value),
                "style": raw_style,
                "abv_source": raw_abv,
                "abv": parse_abv(raw_abv),
                "special_ingredients": raw_special,
                "promotion_allowed": clean(sheet.cell(row_no, 12).value),
                "work_code": work_code,
                "short_code": short_code,
                "uuid": stable_uuid(row_no, work_code),
                "label_code": f"BE-IMP{COMPETITION_ID}-{row_no:03d}",
                "scan_token": stable_token(row_no, work_code),
            }
        )

    if len(rows) != 179:
        raise ValueError(f"已付款记录数量应为 179，实际为 {len(rows)}")
    if len({row["short_code"] for row in rows}) != len(rows):
        raise ValueError("后五位作品编号存在重复")
    if len({row["uuid"] for row in rows}) != len(rows):
        raise ValueError("生成的酒款 UUID 存在重复")
    if len({row["label_code"] for row in rows}) != len(rows):
        raise ValueError("生成的标签编码存在重复")
    if len({row["scan_token"] for row in rows}) != len(rows):
        raise ValueError("生成的扫码 token 存在重复")
    return rows


def build_breweries(rows: list[dict]) -> OrderedDict[str, dict]:
    breweries: OrderedDict[str, dict] = OrderedDict()
    for row in rows:
        item = breweries.setdefault(
            row["brand"],
            {
                "company_name": row["brand"],
                "contact_name": row["contact_name"],
                "wechat": row["wechat"],
                "phones": [],
                "source_brands": [],
                "contact_names": [],
                "source_rows": [],
            },
        )
        if row["phone"] and row["phone"] not in item["phones"]:
            item["phones"].append(row["phone"])
        if row["raw_brand"] not in item["source_brands"]:
            item["source_brands"].append(row["raw_brand"])
        if row["contact_name"] and row["contact_name"] not in item["contact_names"]:
            item["contact_names"].append(row["contact_name"])
        item["source_rows"].append(row["source_row"])
        if not item["contact_name"] and row["contact_name"]:
            item["contact_name"] = row["contact_name"]
        if not item["wechat"] and row["wechat"]:
            item["wechat"] = row["wechat"]
    return breweries


def validate_phone_collisions(rows: list[dict]) -> None:
    phone_brands: defaultdict[str, set[str]] = defaultdict(set)
    for row in rows:
        if row["phone"]:
            phone_brands[row["phone"]].add(row["brand"])
    unresolved = {phone: sorted(brands) for phone, brands in phone_brands.items() if len(brands) > 1}
    if unresolved:
        raise ValueError(f"合并后仍存在手机号对应多个厂牌: {unresolved}")


def write_preview(rows: list[dict], breweries: OrderedDict[str, dict], output_dir: Path) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    report_path = output_dir / "beer_import_preview_20260801.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "导入预览"
    headers = [
        "Excel行号", "联系人", "微信号", "手机号", "原厂牌名", "导入厂牌名", "报名组别",
        "酒款名称", "基础风格（原样）", "ABV原文", "ABV数值", "特殊原料/工艺", "允许宣传",
        "原作品编号", "五位作品编号", "账号处理", "备注",
    ]
    sheet.append(headers)
    no_phone_brands = {name for name, info in breweries.items() if not info["phones"]}
    for row in rows:
        note = []
        if row["source_row"] == 81:
            note.append("基础风格按原表保留：12.3")
        if row["source_row"] == 139:
            note.append("已按确认交换基础风格和ABV")
        if row["raw_brand"] != row["brand"]:
            note.append(f"原厂牌并入{row['brand']}")
        if row["brand"] in no_phone_brands:
            note.append("无手机号，暂不创建账号")
        elif row["phone"]:
            note.append("按手机号预创建或复用账号")
        sheet.append(
            [
                row["source_row"], row["contact_name"], row["wechat"], row["phone"], row["raw_brand"],
                row["brand"], row["group"], row["name"], row["style"], row["abv_source"], float(row["abv"]),
                row["special_ingredients"], row["promotion_allowed"], row["work_code"], row["short_code"],
                "不创建账号" if not row["phone"] else "预创建/复用", "；".join(note),
            ]
        )
    for column in sheet.columns:
        width = min(max(len(str(cell.value or "")) for cell in column) + 2, 42)
        sheet.column_dimensions[column[0].column_letter].width = width
    sheet.freeze_panes = "A2"

    summary = workbook.create_sheet("导入摘要")
    summary.append(["项目", "值"])
    summary_rows = [
        ("比赛ID", COMPETITION_ID),
        ("已付款酒款", len(rows)),
        ("最终厂牌数", len(breweries)),
        ("有手机号酒款", sum(bool(row["phone"]) for row in rows)),
        ("无手机号厂牌", sum(not info["phones"] for info in breweries.values())),
        ("不同基础风格", len({row["style"] for row in rows})),
        ("作品编号重复", "否"),
        ("风格策略", "原始名称保留，写入比赛专用风格配置"),
        ("账号策略", "已知手机号预创建账号，无手机号不创建"),
    ]
    for item in summary_rows:
        summary.append(item)
    workbook.save(report_path)

    csv_path = output_dir / "beer_import_preview_20260801.csv"
    with csv_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(headers)
        for row in rows:
            writer.writerow([
                row["source_row"], row["contact_name"], row["wechat"], row["phone"], row["raw_brand"],
                row["brand"], row["group"], row["name"], row["style"], row["abv_source"], row["abv"],
                row["special_ingredients"], row["promotion_allowed"], row["work_code"], row["short_code"],
                "不创建账号" if not row["phone"] else "预创建/复用", "",
            ])


def write_sql(rows: list[dict], breweries: OrderedDict[str, dict], output_dir: Path) -> tuple[Path, Path, Path, Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    import_path = output_dir / "import_beer_competition_1070_20260801.sql"
    rollback_path = output_dir / "rollback_beer_competition_1070_20260801.sql"
    check_path = output_dir / "check_beer_competition_1070_20260801.sql"
    repair_extra_fields_path = output_dir / "repair_extra_fields_1070_20260801.sql"

    styles = OrderedDict()
    for row in rows:
        styles.setdefault(row["style"], len(styles) + 1)

    lines = [
        "-- 首届中国拉格大赛报名表导入（比赛 1070）",
        "-- 生成时间：2026-08-01",
        "-- 规则：已付款 179 条；基础风格原样保留；有手机号预创建账号。",
        "SET NAMES utf8mb4;",
        "SET time_zone = '+08:00';",
        "START TRANSACTION;",
        "",
        "-- 1. 比赛专用风格配置：不修改全局风格库。",
    ]
    style_vars = {}
    for style, order in styles.items():
        var = f"@style_{order:03d}"
        style_vars[style] = var
        code = f"IMPORT-{COMPETITION_ID}-{order:03d}"
        lines.extend(
            [
                f"SET {var} = (SELECT id FROM competition_style_config WHERE competition_id = {COMPETITION_ID} AND name = {sql_text(style)} ORDER BY id DESC LIMIT 1);",
                f"INSERT INTO competition_style_config (competition_id, name, category_name, style_code, description, sort_order, active_flag, source_library_version) "
                f"SELECT {COMPETITION_ID}, {sql_text(style)}, {sql_text('报名表原始风格')}, {sql_text(code)}, {sql_text('来源：首届中国拉格大赛报名表，厂商原始填写内容')}, {order}, 1, {sql_text(SOURCE_VERSION)} "
                f"WHERE {var} IS NULL;",
                f"SET {var} = COALESCE({var}, LAST_INSERT_ID());",
            ]
        )

    brewery_vars = {}
    new_account_phones = []
    lines.append("-- 2. 厂牌和账号：同手机号厂牌已合并；一个厂牌可以有多个手机号账号。")
    for index, (brand, info) in enumerate(breweries.items(), start=1):
        var = f"@brewery_{index:03d}"
        brewery_vars[brand] = var
        phone_values = [sql_text(phone) for phone in info["phones"]]
        phone_condition = ", ".join(phone_values) if phone_values else sql_text("")
        if info["phones"]:
            lines.append(
                f"SET {var} = (SELECT MIN(b.id) FROM brewery b LEFT JOIN portal_account a ON a.brewery_id = b.id "
                f"WHERE b.phone IN ({phone_condition}) OR a.phone IN ({phone_condition}));"
            )
        else:
            lines.append(
                f"SET {var} = (SELECT id FROM brewery WHERE company_name = {sql_text(brand)} AND phone = {sql_text('')} ORDER BY id DESC LIMIT 1);"
            )
        lines.append(
            f"INSERT INTO brewery (company_name, contact_name, phone, wechat) "
            f"SELECT {sql_text(brand)}, {sql_text(info['contact_name'] or '待完善')}, {sql_text(info['phones'][0] if info['phones'] else '')}, {sql_text(info['wechat']) if info['wechat'] else 'NULL'} "
            f"WHERE {var} IS NULL;"
        )
        lines.append(f"SET {var} = COALESCE({var}, LAST_INSERT_ID());")
        if info["phones"]:
            lines.append(
                f"UPDATE brewery b SET company_name = {sql_text(brand)}, "
                f"contact_name = {sql_text(info['contact_name'] or '待完善')}, "
                f"phone = {sql_text(info['phones'][0])}, "
                f"wechat = {sql_text(info['wechat']) if info['wechat'] else 'NULL'} "
                f"WHERE b.id = {var} AND NOT EXISTS ("
                f"SELECT 1 FROM beer_entry e WHERE e.brewery_id = b.id AND e.deleted_flag = 0"
                f");"
            )
        for phone in info["phones"]:
            new_account_phones.append(phone)
            lines.extend(
                [
                    f"INSERT INTO portal_account (phone, wechat, display_name, brewery_id, status) "
                    f"SELECT {sql_text(phone)}, {sql_text(info['wechat']) if info['wechat'] else 'NULL'}, {sql_text(brand)}, {var}, 1 "
                    f"WHERE NOT EXISTS (SELECT 1 FROM portal_account WHERE phone = {sql_text(phone)});",
                    f"UPDATE portal_account SET brewery_id = {var}, "
                    f"display_name = CASE WHEN display_name IS NULL OR display_name = '' "
                    f"OR display_name LIKE {sql_text('待完善厂牌%')} THEN {sql_text(brand)} ELSE display_name END, "
                    f"wechat = CASE WHEN wechat IS NULL OR wechat = '' THEN "
                    f"{sql_text(info['wechat']) if info['wechat'] else 'NULL'} ELSE wechat END "
                    f"WHERE phone = {sql_text(phone)};",
                ]
            )

    lines.append("-- 3. 酒款和现场标签。REGISTERED 在系统中按已付款状态展示。")
    entry_vars = []
    for row in rows:
        entry_var = f"@entry_{row['source_row']:03d}"
        entry_vars.append((row["uuid"], row["source_row"]))
        extra = {
            "specialIngredients": row["special_ingredients"],
            "custom_1785549677921_0089y": row["promotion_allowed"],
        }
        extra_json = json.dumps(extra, ensure_ascii=False, separators=(",", ":"))
        lines.extend(
            [
                f"INSERT INTO beer_entry (uuid, competition_id, brewery_id, registration_batch_id, category_id, name, style, style_config_id, abv, extra_fields_json, status, stored_flag, deleted_flag) "
                f"VALUES ({sql_text(row['uuid'])}, {COMPETITION_ID}, {brewery_vars[row['brand']]}, NULL, {row['category_id']}, {sql_text(row['name'])}, {sql_text(row['style'])}, {style_vars[row['style']]}, {sql_number(row['abv'])}, {sql_text(extra_json)}, 'REGISTERED', 0, 0);",
                f"SET {entry_var} = LAST_INSERT_ID();",
            ]
        )
        for field_key, field_value in extra.items():
            lines.append(
                f"INSERT INTO beer_entry_extra_field (beer_entry_id, field_key, field_label, field_value) "
                f"SELECT {entry_var}, c.field_key, c.field_label, {sql_text(field_value)} "
                f"FROM entry_field_config c WHERE c.competition_id = {COMPETITION_ID} "
                f"AND c.field_key = {sql_text(field_key)} AND c.active_flag = 1 LIMIT 1;"
            )
        lines.extend(
            [
                f"INSERT INTO entry_scan_label (competition_id, beer_entry_id, label_code, short_code, scan_token, status, generated_by, generated_time) "
                f"VALUES ({COMPETITION_ID}, {entry_var}, {sql_text(row['label_code'])}, {sql_text(row['short_code'])}, {sql_text(row['scan_token'])}, 'ACTIVE', NULL, NOW());",
            ]
        )
    lines.extend(
        [
            "COMMIT;",
            "",
            "-- 导入完成后执行 check_beer_competition_1070_20260801.sql 验收。",
        ]
    )
    import_path.write_text("\n".join(lines) + "\n", encoding="utf-8")

    uuids = ",\n  ".join(sql_text(value) for value, _ in entry_vars)
    new_phones = sorted(set(new_account_phones))
    rollback_lines = [
        "-- 首届中国拉格大赛报名表导入回滚脚本。执行前确认没有在本批酒款上产生评审、送样或成绩数据。",
        "SET NAMES utf8mb4;",
        "START TRANSACTION;",
        "DELETE FROM beer_entry_extra_field WHERE beer_entry_id IN (SELECT id FROM beer_entry WHERE uuid IN (\n  " + uuids + "\n));",
        "DELETE FROM entry_scan_label WHERE beer_entry_id IN (SELECT id FROM beer_entry WHERE uuid IN (\n  " + uuids + "\n));",
        "DELETE FROM beer_entry WHERE uuid IN (\n  " + uuids + "\n);",
        f"DELETE FROM competition_style_config WHERE competition_id = {COMPETITION_ID} AND source_library_version = {sql_text(SOURCE_VERSION)};",
    ]
    if new_phones:
        rollback_lines.append(
            "DELETE FROM portal_account WHERE phone IN (" + ", ".join(sql_text(p) for p in new_phones) + ") "
            "AND id NOT IN (SELECT id FROM (SELECT MIN(id) AS id FROM portal_account WHERE phone IN (" + ", ".join(sql_text(p) for p in new_phones) + ") GROUP BY phone) keep_accounts);"
        )
    rollback_lines.extend(
        [
            "-- 新厂牌 ID 由导入时动态生成，建议优先使用导入前备份恢复；此脚本不删除可能已被其他数据引用的厂牌。",
            "COMMIT;",
        ]
    )
    rollback_path.write_text("\n".join(rollback_lines) + "\n", encoding="utf-8")

    repair_extra_field_lines = [
        "-- 修复比赛 1070 本批导入酒款的补充字段明细。可重复执行，不修改酒款 JSON 快照。",
        "SET NAMES utf8mb4;",
        "SET time_zone = '+08:00';",
        "START TRANSACTION;",
        "",
        "UPDATE beer_entry_extra_field x",
        "JOIN beer_entry e ON e.id = x.beer_entry_id",
        "JOIN entry_scan_label l ON l.beer_entry_id = e.id",
        "  AND l.competition_id = 1070",
        "  AND l.label_code LIKE 'BE-IMP1070-%'",
        "JOIN entry_field_config c ON c.competition_id = e.competition_id",
        "  AND c.field_key = x.field_key",
        "  AND c.active_flag = 1",
        "SET x.field_label = c.field_label,",
        "    x.field_value = JSON_UNQUOTE(JSON_EXTRACT(e.extra_fields_json, CONCAT('$.', c.field_key)))",
        "WHERE e.competition_id = 1070",
        "  AND e.deleted_flag = 0",
        "  AND JSON_VALID(e.extra_fields_json) = 1",
        "  AND JSON_CONTAINS_PATH(e.extra_fields_json, 'one', CONCAT('$.', c.field_key)) = 1",
        "  AND NULLIF(JSON_UNQUOTE(JSON_EXTRACT(e.extra_fields_json, CONCAT('$.', c.field_key))), '') IS NOT NULL;",
        "",
        "INSERT INTO beer_entry_extra_field (beer_entry_id, field_key, field_label, field_value)",
        "SELECT e.id,",
        "       c.field_key,",
        "       c.field_label,",
        "       JSON_UNQUOTE(JSON_EXTRACT(e.extra_fields_json, CONCAT('$.', c.field_key)))",
        "FROM beer_entry e",
        "JOIN entry_scan_label l ON l.beer_entry_id = e.id",
        "  AND l.competition_id = 1070",
        "  AND l.label_code LIKE 'BE-IMP1070-%'",
        "JOIN entry_field_config c ON c.competition_id = e.competition_id",
        "  AND c.active_flag = 1",
        "WHERE e.competition_id = 1070",
        "  AND e.deleted_flag = 0",
        "  AND JSON_VALID(e.extra_fields_json) = 1",
        "  AND JSON_CONTAINS_PATH(e.extra_fields_json, 'one', CONCAT('$.', c.field_key)) = 1",
        "  AND NULLIF(JSON_UNQUOTE(JSON_EXTRACT(e.extra_fields_json, CONCAT('$.', c.field_key))), '') IS NOT NULL",
        "  AND NOT EXISTS (",
        "      SELECT 1",
        "      FROM beer_entry_extra_field existing",
        "      WHERE existing.beer_entry_id = e.id",
        "        AND existing.field_key = c.field_key",
        "  );",
        "",
        "COMMIT;",
        "",
        "SELECT x.field_key, COUNT(*) AS value_count, COUNT(DISTINCT x.beer_entry_id) AS entry_count",
        "FROM beer_entry_extra_field x",
        "JOIN beer_entry e ON e.id = x.beer_entry_id",
        "JOIN entry_scan_label l ON l.beer_entry_id = e.id",
        "WHERE e.competition_id = 1070",
        "  AND e.deleted_flag = 0",
        "  AND l.competition_id = 1070",
        "  AND l.label_code LIKE 'BE-IMP1070-%'",
        "GROUP BY x.field_key",
        "ORDER BY x.field_key;",
    ]
    repair_extra_fields_path.write_text("\n".join(repair_extra_field_lines) + "\n", encoding="utf-8")

    check_lines = [
        "SET NAMES utf8mb4;",
        f"SELECT 'competition_entries', COUNT(*) AS value FROM beer_entry WHERE competition_id = {COMPETITION_ID} AND deleted_flag = 0;",
        f"SELECT 'competition_labels', COUNT(*) AS value FROM entry_scan_label WHERE competition_id = {COMPETITION_ID} AND status = 'ACTIVE';",
        f"SELECT 'competition_styles', COUNT(*) AS value FROM competition_style_config WHERE competition_id = {COMPETITION_ID} AND source_library_version = {sql_text(SOURCE_VERSION)};",
        f"SELECT 'portal_accounts_for_competition', COUNT(DISTINCT a.id) AS value FROM portal_account a JOIN beer_entry e ON e.brewery_id = a.brewery_id WHERE e.competition_id = {COMPETITION_ID} AND e.deleted_flag = 0;",
        f"SELECT 'short_code_not_five', COUNT(*) AS value FROM entry_scan_label WHERE competition_id = {COMPETITION_ID} AND (CHAR_LENGTH(short_code) <> 5 OR short_code NOT REGEXP '^[0-9]{{5}}$');",
        f"SELECT 'short_code_duplicates', COUNT(*) - COUNT(DISTINCT short_code) AS value FROM entry_scan_label WHERE competition_id = {COMPETITION_ID};",
        f"SELECT 'normalized_extra_fields', COUNT(*) AS value FROM beer_entry_extra_field x JOIN beer_entry e ON e.id = x.beer_entry_id WHERE e.competition_id = {COMPETITION_ID} AND e.deleted_flag = 0;",
        f"SELECT x.field_key, COUNT(*) AS value FROM beer_entry_extra_field x JOIN beer_entry e ON e.id = x.beer_entry_id WHERE e.competition_id = {COMPETITION_ID} AND e.deleted_flag = 0 GROUP BY x.field_key ORDER BY x.field_key;",
        f"SELECT e.name, e.style, e.abv, l.short_code FROM beer_entry e JOIN entry_scan_label l ON l.beer_entry_id = e.id WHERE e.competition_id = {COMPETITION_ID} AND e.uuid IN ({uuids}) ORDER BY e.id LIMIT 10;",
        f"SELECT b.company_name, COUNT(*) AS entries, COUNT(DISTINCT a.id) AS accounts FROM beer_entry e JOIN brewery b ON b.id = e.brewery_id LEFT JOIN portal_account a ON a.brewery_id = b.id WHERE e.competition_id = {COMPETITION_ID} AND e.deleted_flag = 0 GROUP BY b.id, b.company_name ORDER BY b.company_name;",
    ]
    check_path.write_text("\n".join(check_lines) + "\n", encoding="utf-8")

    manifest = {
        "competition_id": COMPETITION_ID,
        "source_version": SOURCE_VERSION,
        "entry_uuids": [value for value, _ in entry_vars],
        "new_account_candidate_phones": sorted(set(new_account_phones)),
        "brewery_names": list(breweries),
        "style_names": list(styles),
    }
    (output_dir / "beer_import_manifest_20260801.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    return import_path, rollback_path, check_path, repair_extra_fields_path


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", type=Path, default=INPUT_PATH)
    parser.add_argument("--output-dir", type=Path, default=Path("docs/sql/20260801_beer_import"))
    args = parser.parse_args()

    rows = load_rows(args.input)
    validate_phone_collisions(rows)
    breweries = build_breweries(rows)
    write_preview(rows, breweries, args.output_dir)
    paths = write_sql(rows, breweries, args.output_dir)
    print(f"rows={len(rows)}")
    print(f"breweries={len(breweries)}")
    print(f"phones={len({row['phone'] for row in rows if row['phone']})}")
    print(f"styles={len({row['style'] for row in rows})}")
    for path in paths:
        print(path.resolve())


if __name__ == "__main__":
    main()
